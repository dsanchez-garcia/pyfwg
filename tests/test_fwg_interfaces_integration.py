"""Real FWG Engine JAR coverage across the supported public interfaces.

Unlike unit tests, these tests launch the supplied Java archives and inspect
their final weather files. No archive or baseline is downloaded automatically.
"""
import csv
import hashlib
import os
from pathlib import Path
import shutil

import pandas as pd
import pytest
from openpyxl import load_workbook

from pyfwg import (
    MorphingIterator, MorphingWorkflowCordex, MorphingWorkflowEurope, MorphingWorkflowGlobal,
    export_template_to_excel, get_fwg_distribution_info, load_runs_from_excel,
    morph_epw_cordex, morph_epw_europe, morph_epw_global,
)
from pyfwg.utils import detect_fwg_tool, detect_fwg_version, get_fwg_cli_style

pytestmark = [pytest.mark.integration, pytest.mark.interfaces]
CLASSES = {"global": MorphingWorkflowGlobal, "europe": MorphingWorkflowEurope, "cordex": MorphingWorkflowCordex}
APIS = {"global": morph_epw_global, "europe": morph_epw_europe, "cordex": morph_epw_cordex}


def discover(config):
    folders = [config.getoption("--fwg-jar-dir"), *config.getoption("--fwg-extra-jar-dir")]
    jars = set()
    for folder in filter(None, folders):
        directory = Path(folder).expanduser().resolve()
        if not directory.is_dir():
            raise pytest.UsageError(f"FWG JAR directory does not exist: {directory}")
        jars.update(jar.resolve() for jar in directory.glob("*.jar")
                    if detect_fwg_tool(str(jar)) in CLASSES
                    and get_fwg_distribution_info(str(jar)) is not None)
    return sorted(jars)


def pytest_generate_tests(metafunc):
    if "installed_jar" in metafunc.fixturenames:
        cases = [pytest.param(jar, id=jar.stem) for jar in discover(metafunc.config)]
        if not cases:
            cases = [pytest.param(None, marks=pytest.mark.skip(reason="Supply --fwg-jar-dir with actual FWG archives"))]
        metafunc.parametrize("installed_jar", cases)


def descriptor(jar):
    info = get_fwg_distribution_info(str(jar))
    tool = detect_fwg_tool(str(jar))
    version = detect_fwg_version(str(jar))
    engine = get_fwg_cli_style(version, tool) == "engine"
    if engine:
        assert info is not None, f"Cannot read the FWG Engine catalog: {jar}"
    return tool, engine, info


def baseline_for(jar, config):
    tool, _, info = descriptor(jar)
    directory = config.getoption("--fwg-epw-dir")
    if not directory:
        pytest.skip("Supply --fwg-epw-dir with real baselines; EUR-11.epw also serves the Global examples")
    domain = (info or {}).get("domain") or "EUR-11"
    epw = Path(directory).expanduser().resolve() / f"{domain}.epw"
    if not epw.is_file():
        pytest.skip(f"Missing real {tool} baseline: {epw}")
    return epw


def options_for(jar, epw, config, *, ensemble=False):
    tool, engine, info = descriptor(jar)
    model_key = "fwg_gcms" if tool == "global" else "fwg_rcm_pairs"
    options = {"fwg_jar_path": str(jar), "delete_temp_files": False,
               "fwg_create_ensemble": ensemble, "fwg_use_multithreading": False}
    if tool != "cordex":
        options["fwg_add_uhi"] = False
    if engine:
        models = info["ensemble_models"] if ensemble else [
            "CanESM5" if tool == "global" and "CanESM5" in info["models"] else info["models"][0]]
        if ensemble and config.getoption("--fwg-ensemble-size") is not None:
            models = models[:config.getoption("--fwg-ensemble-size")]
        assert models and (not ensemble or len(models) >= 2)
        details = info["model_details"]
        scenario = next(s for s in info["scenarios"] if all(s in details[m]["scenarios"] for m in models))
        period = next(t for t in info["timeframes"] if all(t in details[m]["timeframes"] for m in models))
        options.update({model_key: models, "fwg_scenarios": [scenario], "fwg_timeframes": [int(period)],
                        "fwg_output_profile": "MINIMAL", "fwg_interpolation_method_id": "IDW",
                        "fwg_extra_args": {"threads": 1, "io_threads": 1, "save_technical_log": True}})
        cache = config.getoption("--fwg-cache-dir")
        if cache:
            options["fwg_model_unzipped_folder"] = str(
                Path(cache).expanduser().resolve() / info["distribution_id"] / info["version"])
    else:
        options[model_key] = ["CanESM5"] if tool == "global" else ["ICHEC_EC_EARTH_SMHI_RCA4"]
        options["fwg_interpolation_method_id"] = 0
    return options


@pytest.fixture
def real_java(request, monkeypatch):
    java = shutil.which(request.config.getoption("--fwg-java"))
    if not java:
        pytest.skip("Java unavailable; supply --fwg-java")
    monkeypatch.setenv("PATH", str(Path(java).parent) + os.pathsep + os.environ.get("PATH", ""))
    return java


def verify_outputs(output, epw, expected_stems=None):
    files = list(output.glob("*"))
    weather = sorted(output.glob("*.epw"))
    stats = sorted(output.glob("*.stat"))
    assert weather and len(weather) == len(stats), f"Missing EPW/STAT output in {output}; inspect preserved diagnostics"
    assert {p.stem for p in weather} == {p.stem for p in stats}
    assert len(files) == 2 * len(weather), f"Unexpected collected artifacts: {files}"
    if expected_stems is not None:
        assert {p.stem for p in weather} == set(expected_stems)
    original = list(csv.reader(epw.read_text(encoding="latin-1").splitlines()))
    for path in weather:
        rows = list(csv.reader(path.read_text(encoding="latin-1").splitlines()))
        assert rows[0][0] == "LOCATION" and len(rows) == len(original) and len(rows) >= 8768
        assert all(abs(float(rows[0][i]) - float(original[0][i])) < 1e-4 for i in (6, 7))
        assert any(a[6] != b[6] for a, b in zip(rows[8:], original[8:])), "Output temperature is unchanged"
    assert all(p.stat().st_size > 0 for p in files)


def excel_round_trip(iterator, rows, path):
    """Fill the real exported template with Excel cell types, then use pyfwg's loader."""
    export_template_to_excel(iterator, str(path))
    book = load_workbook(path)
    sheet = book.active
    columns = {cell.value: cell.column for cell in sheet[1]}
    for index, row in enumerate(rows, start=2):
        for key, value in row.items():
            assert key in columns, f"Exported template lacks {key}"
            sheet.cell(index, columns[key], repr(value) if isinstance(value, (list, dict)) else value)
    book.save(path)
    book.close()
    loaded = load_runs_from_excel(str(path))
    assert len(loaded) == len(rows)
    for index, row in enumerate(rows):
        for key, value in row.items():
            if isinstance(value, (list, dict)):
                assert loaded.loc[index, key] == value, f"Excel conversion changed {key}"
    return loaded


@pytest.mark.parametrize("workflow_class", list(CLASSES.values()), ids=lambda cls: cls.__name__)
def test_installed_jar_class_compatibility(installed_jar, workflow_class, tmp_path):
    tool, engine, info = descriptor(installed_jar)
    expected = workflow_class is CLASSES[tool] or (tool == "europe" and engine and workflow_class is MorphingWorkflowCordex)
    epw = tmp_path / "station.epw"
    epw.write_text("Configuration check only")
    wf = workflow_class()
    wf.map_categories([str(epw)], keyword_mapping={"station": {"Station": "station"}})
    wf.configure_and_preview(
        fwg_jar_path=str(installed_jar), final_output_dir=str(tmp_path / "out"),
        temp_base_dir=str(tmp_path / "temp"),
        output_filename_pattern=f"{{station}}_{{{wf.scenario_placeholder_name}}}_{{year}}",
    )
    assert wf.is_config_valid is expected


@pytest.mark.java
@pytest.mark.morphing
@pytest.mark.parametrize("route", ["api", "workflow", "iterator-dataframe", "iterator-xlsx"])
def test_installed_jar_public_interfaces(installed_jar, route, tmp_path, request, real_java):
    tool, engine, info = descriptor(installed_jar)
    epw = baseline_for(installed_jar, request.config)
    original_sha = hashlib.sha256(epw.read_bytes()).hexdigest()
    options = options_for(installed_jar, epw, request.config)
    output = tmp_path / "out"
    temp = tmp_path / "temp"
    options["temp_base_dir"] = str(temp)
    cls = CLASSES[tool]
    pattern = f"{{station}}_{{{cls.scenario_placeholder_name}}}_{{year}}"
    mapping = {"station": {"Station": epw.stem}}
    expected = None
    if route == "api":
        returned = APIS[tool](epw_paths=[str(epw)], output_dir=str(output), **options)
        assert returned and set(map(Path, returned)) == set(output.glob("*"))
        if engine:
            assert len(returned) == 2
    elif route == "workflow":
        wf = cls()
        wf.map_categories([str(epw)], keyword_mapping=mapping)
        wf.configure_and_preview(final_output_dir=str(output), output_filename_pattern=pattern, **options)
        assert wf.is_config_valid
        expected = [Path(p).stem for p in wf.rename_plan[str(epw)].values()]
        wf.execute_morphing()
    else:
        iterator = MorphingIterator(cls)
        iterator.set_default_values(output_filename_pattern=pattern)
        row = {**options, "epw_paths": [str(epw)], "final_output_dir": str(output), "keyword_mapping": mapping,
               "scenario_mapping": {s: s.upper() for s in (info["scenarios"] if info else cls.tool_scenarios)}}
        # Exercise explicit overrides of nested fwg_params and integer cells read as numbers.
        row["fwg_params"] = {"month_transition_hours": 48, "create_ensemble": True}
        row["fwg_month_transition_hours"] = 24
        runs = excel_round_trip(iterator, [row], tmp_path / "runs.xlsx") if route == "iterator-xlsx" else pd.DataFrame([row])
        iterator.generate_morphing_workflows(runs)
        assert len(iterator.prepared_workflows) == 1
        wf = iterator.prepared_workflows[0]
        assert wf.is_config_valid
        assert wf.inputs["fwg_params"]["month_transition_hours"] == 24
        assert wf.inputs["fwg_params"]["create_ensemble"] is False
        expected = [Path(p).stem for p in wf.rename_plan[str(epw)].values()]
        iterator.run_morphing_workflows()
    verify_outputs(output, epw, expected)
    assert hashlib.sha256(epw.read_bytes()).hexdigest() == original_sha


@pytest.mark.java
@pytest.mark.morphing
@pytest.mark.parametrize("ensemble", [False, True], ids=["single-model", "full-ensemble"])
def test_cordex_mixed_regions_xlsx_batch(tmp_path, request, real_java, ensemble):
    jars = [jar for jar in discover(request.config) if (get_fwg_distribution_info(str(jar)) or {}).get("framework") == "CORDEX"]
    if len(jars) < 2:
        pytest.skip("Supply at least two CORDEX Engine JARs to exercise a mixed-region XLSX batch")
    iterator = MorphingIterator(MorphingWorkflowCordex)
    iterator.set_default_values(output_filename_pattern="{region}_{rcp}_{year}")
    rows, expected_baselines = [], []
    for index, jar in enumerate(jars):
        epw = baseline_for(jar, request.config)
        info = get_fwg_distribution_info(str(jar))
        domain = info["domain"]
        options = options_for(jar, epw, request.config, ensemble=ensemble)
        options.pop("fwg_add_uhi", None)  # CORDEX does not expose UHI, including its EUR-11 alias.
        # Include blank cells and nested defaults in a real multirow worksheet.
        options.pop("fwg_create_ensemble")
        rows.append({**options, "epw_paths": [str(epw)], "final_output_dir": str(tmp_path / f"out-{index}-{domain}"),
                     "temp_base_dir": str(tmp_path / f"temp-{index}-{domain}"),
                     "keyword_mapping": {"region": {domain: domain}},
                     "fwg_params": {"create_ensemble": ensemble, "month_transition_hours": 48},
                     "fwg_month_transition_hours": 24 if index % 2 else None})
        expected_baselines.append(epw)
    runs = excel_round_trip(iterator, rows, tmp_path / "mixed-regions.xlsx")
    iterator.generate_morphing_workflows(runs)
    assert len(iterator.prepared_workflows) == len(rows)
    for index, (wf, row) in enumerate(zip(iterator.prepared_workflows, rows)):
        assert wf.is_config_valid
        assert wf.cordex_domain in row["final_output_dir"]
        assert wf.inputs["fwg_params"]["month_transition_hours"] == (24 if index % 2 else 48)
        assert wf.inputs["fwg_params"]["create_ensemble"] is ensemble
    iterator.run_morphing_workflows()
    for wf, row, epw in zip(iterator.prepared_workflows, rows, expected_baselines):
        expected = [Path(p).stem for p in wf.rename_plan[str(epw)].values()]
        verify_outputs(Path(row["final_output_dir"]), epw, expected)
